import re
import openpyxl
from openpyxl.styles import Font, PatternFill

content = open('lib/data/content/wiring_content.dart', 'r').read()
pattern = r"WiringDiagram\([^)]*id:\s*'([^']+)'[^)]*title:\s*'([^']+)'"
matches = re.findall(pattern, content)

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Diagram Review Record"

headers = [
    'Diagram ID', 'Diagram title', 'Intended country or standard', 'Supply type',
    'Voltage and phase', 'Earthing arrangement', 'Source reference', 'Reviewer name',
    'Reviewer qualification', 'Review date', 'Revision number', 
    'Status: Draft, Reviewed, Approved or Rejected', 'Required corrections'
]

# Write headers
for col_num, header_title in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col_num, value=header_title)
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")

# Write data
for row_num, match in enumerate(matches, 2):
    d_id = match[0]
    d_title = match[1].replace("\\'", "'")
    row_data = [d_id, d_title, 'Generic/Global', 'AC', '', '', '', '', '', '', '1.0', 'Draft', '']
    for col_num, value in enumerate(row_data, 1):
        ws.cell(row=row_num, column=col_num, value=value)

# Auto-size columns
for col in ws.columns:
    max_length = 0
    column = col[0].column_letter
    for cell in col:
        try:
            if len(str(cell.value)) > max_length:
                max_length = len(str(cell.value))
        except:
            pass
    adjusted_width = (max_length + 2)
    ws.column_dimensions[column].width = adjusted_width

wb.save('DIAGRAM_REVIEW_RECORD.xlsx')
print("Excel file created.")
